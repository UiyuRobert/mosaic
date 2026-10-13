#!/usr/bin/env python3
"""Real-collection validation of the aero-engine starter keyword list.

Local acceptance harness — reads the Excel workbook (起步采集清单 joined with
术语词表 and 检索任务), runs tiered real searches across free sources,
scores topical relevance, downloads legal OA PDFs through the MOSAIC queue,
verifies file + identity, and records full evidence.

This is a *test harness*, not a product feature: MOSAIC stays a generic
keyword search tool. State lives under ``--output-dir`` (never committed).

Usage:
    python tools/validate_starter_list.py \
        --xlsx "/path/航发设计_论文采集中英关键词与检索图谱.xlsx" \
        --output-dir ./local-validation --resume
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

log = logging.getLogger("validate")

# ---------------------------------------------------------------------------
# Workbook parsing
# ---------------------------------------------------------------------------

EXPECTED = {"rows": 321, "domains": 17, "batches": {"P0": 182, "P1": 130, "P2": 9}}
BROAD_OK_DOMAINS = {"基础学科", "仿真设计与优化"}  # MEDIUM grade acceptable
BATCH_ORDER = {"P0": 0, "P1": 1, "P2": 2}

# Context vocabulary marking aero-engine / gas-turbine relevance. Used both
# for query construction and relevance scoring.
AERO_CONTEXT = {
    "gas turbine",
    "turbine",
    "turbojet",
    "turbofan",
    "turboprop",
    "turboshaft",
    "aeroengine",
    "aero engine",
    "aircraft engine",
    "jet engine",
    "aero gas turbine",
    "turbomachinery",
    "turbine blade",
    "compressor",
    "combustor",
    "afterburner",
    "nozzle",
    "intake",
    "aircraft",
    "propulsion",
    "航空发动机",
    "燃气轮机",
    "涡轮",
    "压气机",
    "燃烧室",
    "喷管",
    "进气道",
    "叶轮机械",
}


@dataclass
class Term:
    term_id: str
    excel_row: int
    domain: str
    topic: str
    zh_core: str
    zh_ext: list[str]
    en_core: str
    en_ext: list[str]
    abbr: str
    related: list[str]
    focus: str
    boundary: str
    batch: str
    en_wide_q: str
    en_focus_q: str
    zh_wide_q: str
    zh_focus_q: str
    ntrs_q: str
    trial_status: str
    feedback: str
    # parsed from the Excel query strings
    term_alts: list[str] = field(default_factory=list)
    ctx_alts: list[str] = field(default_factory=list)


def _first_line(val) -> str:
    return re.sub(r"\s+", " ", str(val or "").split("\n")[0]).strip()


def _ext_lines(val) -> list[str]:
    """Parse 'core + kuozhan a/b' style cells -> [core, ext_a, ext_b]."""
    if not val:
        return []
    parts = re.split(r"[\n；;]", str(val))
    out = []
    for p in parts:
        p = re.sub(r"^扩展[:：]\s*", "", p.strip())
        for x in re.split(r"[、,，/]", p):
            x = x.strip()
            if x and x not in out:
                out.append(x)
    return out


def _parse_groups(q: str) -> tuple[list[str], list[str]]:
    """Parse '("a" OR "b") AND ("c" OR "d")' → (term_alts, ctx_alts)."""
    if not q:
        return [], []
    groups = re.findall(r"\(([^)]+)\)", q)
    alts = [
        [m.strip().strip('"').strip("'") for m in re.split(r"\s+OR\s+|\|", g, flags=re.I)]
        for g in groups
    ]
    alts = [[a for a in g if a] for g in alts]
    if len(alts) >= 2:
        return alts[0], alts[1]
    if len(alts) == 1:
        return alts[0], []
    return [_first_line(q).strip('"')] if q.strip() else [], []


def load_terms(xlsx: Path) -> list[Term]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws_list, ws_terms, ws_tasks = wb["起步采集清单"], wb["术语词表"], wb["检索任务"]

    tmap: dict[tuple[str, str], dict] = {}
    for r in ws_terms.iter_rows(min_row=7, values_only=True):
        if not r[0]:
            continue
        key = (_first_line(r[3]).lower(), _first_line(r[4]).lower())
        assert key not in tmap, f"duplicate term key {key}"
        tmap[key] = {
            "term_id": r[0],
            "domain": r[1],
            "topic": r[2],
            "zh": r[3],
            "en": r[4],
            "en_ext": r[5],
            "abbr": r[6],
            "zh_ext": r[7],
            "batch": r[8],
            "boundary": r[9] or "",
            "ctx_id": r[13],
        }

    tasks: dict[str, dict] = {}
    for r in ws_tasks.iter_rows(min_row=7, values_only=True):
        if not r[0]:
            continue
        tasks[r[0]] = {
            "batch": r[3],
            "en_wide": r[4] or "",
            "en_focus": r[5] or "",
            "zh_wide": r[6] or "",
            "zh_focus": r[7] or "",
            "ntrs": r[8] or "",
            "trial": r[9] or "",
            "feedback": r[10] or "",
        }

    terms: list[Term] = []
    problems: list[str] = []
    for excel_row, r in enumerate(ws_list.iter_rows(min_row=6, values_only=True), start=6):
        if not any(r):
            continue
        key = (_first_line(r[1]).lower(), _first_line(r[2]).lower())
        meta = tmap.get(key)
        if meta is None:
            problems.append(f"row {excel_row}: no term match for {key}")
            continue
        task = tasks.get(meta["term_id"], {})
        term_alts, ctx_alts = _parse_groups(task.get("en_focus", ""))
        if not term_alts:
            term_alts = [_first_line(r[2])]
        zh_ext = _ext_lines(meta["zh_ext"]) or _ext_lines(r[1])[1:]
        en_ext = _ext_lines(meta["en_ext"]) or _ext_lines(r[2])[1:]
        terms.append(
            Term(
                term_id=meta["term_id"],
                excel_row=excel_row,
                domain=meta["domain"] or _first_line(r[0]),
                topic=meta["topic"] or "",
                zh_core=meta["zh"] or _first_line(r[1]),
                zh_ext=zh_ext,
                en_core=meta["en"] or _first_line(r[2]),
                en_ext=en_ext,
                abbr=_first_line(meta["abbr"]),
                related=_ext_lines(r[3]),
                focus=_first_line(r[4]),
                boundary=str(meta["boundary"] or ""),
                batch=task.get("batch") or meta["batch"] or "P2",
                en_wide_q=task.get("en_wide", ""),
                en_focus_q=task.get("en_focus", ""),
                zh_wide_q=task.get("zh_wide", ""),
                zh_focus_q=task.get("zh_focus", ""),
                ntrs_q=task.get("ntrs", ""),
                trial_status=task.get("trial", ""),
                feedback=task.get("feedback", ""),
                term_alts=term_alts,
                ctx_alts=ctx_alts,
            )
        )

    # Structural validation (warn, not hard-fail, so workbook edits don't break the tool)
    counts = {
        "rows": len(terms),
        "ids": len({t.term_id for t in terms}),
        "domains": len({t.domain for t in terms}),
    }
    batches = {}
    for t in terms:
        batches[t.batch] = batches.get(t.batch, 0) + 1
    log.info(
        "workbook: %d rows, %d unique ids, %d domains, batches=%s",
        counts["rows"],
        counts["ids"],
        counts["domains"],
        batches,
    )
    if counts["rows"] != EXPECTED["rows"] or counts["ids"] != EXPECTED["rows"]:
        log.warning("row/id count drifted from expected %s: %s", EXPECTED, counts)
    for p in problems:
        log.error("JOIN PROBLEM: %s", p)
    terms.sort(key=lambda t: (BATCH_ORDER.get(t.batch, 9), t.term_id))
    return terms


# ---------------------------------------------------------------------------
# Query planning (tiered: focused → extensions → wide → targeted)
# ---------------------------------------------------------------------------


def plan_queries(term: Term, max_queries: int) -> list[tuple[str, str]]:
    """Return [(query, tier)] — deduplicated, ordered by tier.

    Plain-text query strings: each source maps them onto its own API — the
    Excel boolean syntax is parsed into alternates, not forwarded verbatim.
    """
    out: list[tuple[str, str]] = []

    def add(q: str, tier: str) -> None:
        q = re.sub(r"\s+", " ", q).strip()
        if q and all(q != x for x, _ in out):
            out.append((q, tier))

    core = term.en_core
    ctx = term.ctx_alts or ["gas turbine", "aircraft engine"]
    # L1 focused: core + each context alternate (max 2)
    for c in ctx[:2]:
        add(f"{core} {c}", "L1_focus")
    for alt in term.term_alts:
        if alt.lower() != core.lower():
            add(f"{alt} {ctx[0]}", "L1_focus")
    # L2 extension terms (with context)
    for ext in term.en_ext[:3]:
        if ext.lower() != core.lower():
            add(f"{ext} {ctx[0]}", "L2_ext")
    if term.abbr:
        add(f"{term.abbr} {ctx[0]}", "L2_abbr")
    # L3 wide
    add(core, "L3_wide")
    for alt in term.term_alts:
        add(alt, "L3_wide")
    # L4 zh + related rescue (rarely needed; low priority in English sources)
    zh_alts, zh_ctx = _parse_groups(term.zh_focus_q)
    if zh_alts:
        add(f"{zh_alts[0]} {(zh_ctx or ['航空发动机'])[0]}", "L4_zh")
    for rel in term.related[:2]:
        add(f"{rel} {ctx[0]}", "L4_related")
    return out[:max_queries]


# ---------------------------------------------------------------------------
# Relevance scoring (deterministic, explainable — no LLM)
# ---------------------------------------------------------------------------

_WORD = re.compile(r"[a-z0-9\-]+|[一-鿿]+")


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "").lower()
    return re.sub(r"[^a-z0-9一-鿿\- ]+", " ", text)


def _has_phrase(text_norm: str, phrase: str) -> bool:
    p = _norm(phrase).strip()
    if not p:
        return False
    return re.search(rf"(?<![a-z0-9]){re.escape(p)}(?![a-z0-9])", text_norm) is not None


def score_paper(
    term: Term, paper, *, has_abstract_source: bool = True
) -> tuple[str, float, list[str]]:
    """Return (grade, score, reasons). HIGH/MEDIUM/LOW/UNCERTAIN."""
    title = _norm(paper.title or "")
    abstract = _norm(paper.abstract or "")
    score = 0.0
    why: list[str] = []

    phrases = [(term.en_core, 4.0, "core")]
    for e in term.en_ext:
        if e and e.lower() != term.en_core.lower():
            phrases.append((e, 2.5, "ext"))
    for p, pts, tag in phrases:
        if _has_phrase(title, p):
            score += pts
            why.append(f"{tag}«{p}»∈title")
        elif _has_phrase(abstract, p):
            score += pts / 2
            why.append(f"{tag}«{p}»∈abs")
    if term.abbr and len(term.abbr) > 1 and _has_phrase(title, term.abbr):
        score += 1.5
        why.append(f"abbr«{term.abbr}»∈title")
    if term.zh_core and term.zh_core in (paper.title or ""):
        score += 3.0
        why.append("zh∈title")
    elif term.zh_core and term.zh_core in (paper.abstract or ""):
        score += 1.0
        why.append("zh∈abs")

    # Word-level fallback: contiguous-phrase matching misses reordered or
    # compound variants ("heat transfer in a rotating ribbed channel" never
    # contains "rotating channel heat transfer"). Count how many content
    # words of the core phrase appear in the title and award partial
    # credit — all words present is strong topical evidence.
    _stop = {
        "and",
        "the",
        "for",
        "with",
        "of",
        "a",
        "an",
        "in",
        "on",
        "to",
        "between",
        "based",
        "using",
        "via",
        "from",
    }
    words = [w for w in _norm(term.en_core).split() if len(w) > 2 and w not in _stop]
    if len(words) >= 2 and not _has_phrase(title, term.en_core):
        hit = sum(1 for w in words if _has_phrase(title, w))
        frac = hit / len(words)
        if frac >= 1.0:
            score += 2.0
            why.append("core-words∈title")
        elif frac >= 0.6:
            score += 1.0
            why.append("core-words~∈title")
        elif hit:
            score += 0.4

    rel_hits = 0.0
    for rel in term.related:
        if _has_phrase(title, rel):
            rel_hits += 0.75
    if rel_hits:
        score += min(1.5, rel_hits)
        why.append("related∈title")

    ctx_hits = [c for c in AERO_CONTEXT if _has_phrase(title, c) or _has_phrase(abstract, c)]
    if ctx_hits:
        score += 1.0
        why.append(f"aero-ctx:{ctx_hits[0]}")

    if (paper.pub_type or "").lower() in (
        "technical memorandum",
        "technical publication",
        "contractor report",
        "contractor/grantee report",
        "special publication",
    ):
        score += 0.25
        why.append("tech-report")

    core_title = _has_phrase(title, term.en_core)
    core_abs = _has_phrase(abstract, term.en_core)
    if core_title or score >= 5.0 or (core_abs and ctx_hits):
        grade = "HIGH"
    elif score >= 3.0 or core_abs or (score >= 2.5 and term.domain in BROAD_OK_DOMAINS):
        grade = "MEDIUM"
    elif score > 0:
        grade = "LOW"
    else:
        grade = "UNCERTAIN" if not paper.abstract else "LOW"
    return grade, round(score, 2), why


# ---------------------------------------------------------------------------
# Evidence store (resume-capable)
# ---------------------------------------------------------------------------

SCHEMA = """
CREATE TABLE IF NOT EXISTS term_results(
  term_id TEXT PRIMARY KEY, excel_row INT, domain TEXT, batch TEXT,
  zh TEXT, en TEXT, status TEXT, queries_run INT DEFAULT 0,
  candidates INT DEFAULT 0, deduped INT DEFAULT 0, relevant INT DEFAULT 0,
  dl_candidates INT DEFAULT 0, dl_ok INT DEFAULT 0, pdf_ok INT DEFAULT 0,
  reason TEXT, suggestion TEXT, updated_at TEXT);
CREATE TABLE IF NOT EXISTS papers(
  uid TEXT PRIMARY KEY, title TEXT, year INT, doi TEXT, ntrs_id TEXT,
  source TEXT, pdf_url TEXT, local_path TEXT, sha256 TEXT,
  verify_status TEXT, identity_status TEXT, json TEXT);
CREATE TABLE IF NOT EXISTS links(
  term_id TEXT, paper_uid TEXT, grade TEXT, score REAL, downloaded INT DEFAULT 0,
  PRIMARY KEY(term_id, paper_uid));
CREATE TABLE IF NOT EXISTS search_attempts(
  id INTEGER PRIMARY KEY AUTOINCREMENT, term_id TEXT, ts TEXT, query TEXT,
  tier TEXT, source TEXT, http_status TEXT, n_results INT, error TEXT);
CREATE TABLE IF NOT EXISTS download_attempts(
  id INTEGER PRIMARY KEY AUTOINCREMENT, term_id TEXT, paper_uid TEXT, ts TEXT,
  backend TEXT, url TEXT, result TEXT, reason TEXT, path TEXT, sha256 TEXT);
CREATE TABLE IF NOT EXISTS source_health(
  source TEXT PRIMARY KEY, consec_errors INT DEFAULT 0,
  cooldown_until REAL DEFAULT 0, last_error TEXT);
CREATE TABLE IF NOT EXISTS run_meta(k TEXT PRIMARY KEY, v TEXT);
"""


class Store:
    def __init__(self, path: Path):
        import threading

        self._lock = threading.RLock()
        path.parent.mkdir(parents=True, exist_ok=True)
        self.con = sqlite3.connect(str(path), timeout=30, check_same_thread=False)
        self.con.row_factory = sqlite3.Row
        self.con.executescript(SCHEMA)
        self.con.commit()

    def q(self, sql: str, args=()) -> list[sqlite3.Row]:
        with self._lock:
            return self.con.execute(sql, args).fetchall()

    def x(self, sql: str, args=()) -> None:
        with self._lock:
            self.con.execute(sql, args)
            self.con.commit()

    def term_status(self, term_id: str) -> str | None:
        r = self.q("SELECT status FROM term_results WHERE term_id=?", (term_id,))
        return r[0]["status"] if r else None

    def save_term_result(self, t: Term, **kw) -> None:
        self.x(
            """INSERT INTO term_results(term_id,excel_row,domain,batch,zh,en,status,
               queries_run,candidates,deduped,relevant,dl_candidates,dl_ok,pdf_ok,
               reason,suggestion,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(term_id) DO UPDATE SET status=excluded.status,
               queries_run=queries_run+excluded.queries_run,
               candidates=excluded.candidates,deduped=excluded.deduped,
               relevant=excluded.relevant,dl_candidates=excluded.dl_candidates,
               dl_ok=excluded.dl_ok,pdf_ok=excluded.pdf_ok,
               reason=excluded.reason,suggestion=excluded.suggestion,
               updated_at=excluded.updated_at""",
            (
                t.term_id,
                t.excel_row,
                t.domain,
                t.batch,
                t.zh_core,
                t.en_core,
                kw.get("status", "PENDING"),
                kw.get("queries", 0),
                kw.get("candidates", 0),
                kw.get("deduped", 0),
                kw.get("relevant", 0),
                kw.get("dl_candidates", 0),
                kw.get("dl_ok", 0),
                kw.get("pdf_ok", 0),
                kw.get("reason", ""),
                kw.get("suggestion", ""),
                datetime.now(UTC).isoformat(),
            ),
        )

    def save_paper(self, paper, local_path="", sha256="", verify="", identity="") -> None:
        self.x(
            """INSERT INTO papers(uid,title,year,doi,ntrs_id,source,pdf_url,local_path,
               sha256,verify_status,identity_status,json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(uid) DO UPDATE SET title=excluded.title,
               doi=COALESCE(excluded.doi,papers.doi),
               local_path=COALESCE(NULLIF(excluded.local_path,''),papers.local_path),
               sha256=COALESCE(NULLIF(excluded.sha256,''),papers.sha256),
               verify_status=COALESCE(NULLIF(excluded.verify_status,''),papers.verify_status),
               identity_status=COALESCE(NULLIF(excluded.identity_status,''),papers.identity_status)""",
            (
                paper.uid,
                paper.title,
                paper.year,
                paper.doi,
                paper.ntrs_id,
                paper.source,
                paper.pdf_url,
                local_path,
                sha256,
                verify,
                identity,
                json.dumps(paper.to_dict(), ensure_ascii=False),
            ),
        )

    def link(self, term_id: str, uid: str, grade: str, score: float, downloaded: bool) -> None:
        self.x(
            """INSERT INTO links(term_id,paper_uid,grade,score,downloaded) VALUES(?,?,?,?,?)
               ON CONFLICT(term_id,paper_uid) DO UPDATE SET grade=excluded.grade,
               score=excluded.score,downloaded=max(links.downloaded,excluded.downloaded)""",
            (term_id, uid, grade, score, int(downloaded)),
        )

    def log_search(self, term_id, query, tier, source, http_status, n, error="") -> None:
        self.x(
            "INSERT INTO search_attempts(term_id,ts,query,tier,source,http_status,n_results,error)"
            " VALUES(?,?,?,?,?,?,?,?)",
            (
                term_id,
                datetime.now(UTC).isoformat(),
                query,
                tier,
                source,
                http_status,
                n,
                error[:500],
            ),
        )

    def log_download(self, term_id, uid, backend, url, result, reason, path, sha256) -> None:
        self.x(
            "INSERT INTO download_attempts(term_id,paper_uid,ts,backend,url,result,reason,path,sha256)"
            " VALUES(?,?,?,?,?,?,?,?,?)",
            (
                term_id,
                uid,
                datetime.now(UTC).isoformat(),
                backend,
                url,
                result,
                reason[:400],
                path,
                sha256,
            ),
        )

    def source_ok(self, name: str) -> bool:
        r = self.q("SELECT cooldown_until FROM source_health WHERE source=?", (name,))
        return not (r and r[0]["cooldown_until"] > time.time())

    def source_success(self, name: str) -> None:
        self.x(
            "INSERT INTO source_health(source,consec_errors,cooldown_until,last_error)"
            " VALUES(?,0,0,'') ON CONFLICT(source) DO UPDATE SET consec_errors=0,last_error=''",
            (name,),
        )

    def source_fail(self, name: str, error: str, *, cooldown: float = 0) -> None:
        r = self.q("SELECT consec_errors FROM source_health WHERE source=?", (name,))
        n = (r[0]["consec_errors"] if r else 0) + 1
        cd = max(cooldown, 900.0) if n >= 3 else cooldown
        until = time.time() + cd if cd else 0
        self.x(
            "INSERT INTO source_health(source,consec_errors,cooldown_until,last_error)"
            " VALUES(?,?,?,?) ON CONFLICT(source) DO UPDATE SET consec_errors=?,"
            " cooldown_until=max(cooldown_until,?), last_error=?",
            (name, n, until, error[:300], n, until, error[:300]),
        )

    def relevant_links(self, term_id: str) -> list[sqlite3.Row]:
        return self.q(
            "SELECT l.*, p.local_path FROM links l JOIN papers p ON p.uid=l.paper_uid"
            " WHERE l.term_id=? AND l.grade IN ('HIGH','MEDIUM') ORDER BY"
            " CASE l.grade WHEN 'HIGH' THEN 0 ELSE 1 END, l.score DESC",
            (term_id,),
        )


def _dl_hint(paper) -> bool:
    """Cheap 'probably resolvable to a PDF' signal used for try-order only."""
    return bool(paper.pdf_url or paper.pdf_candidates or paper.is_open_access)


# ---------------------------------------------------------------------------
# PDF identity check
# ---------------------------------------------------------------------------


def pdf_identity(path: Path, paper) -> tuple[str, str]:
    """Compare embedded first-page text with the catalogue identity.

    Returns (status, detail): ok|weak|review|mismatch.
    """
    try:
        import fitz
    except ImportError:
        import pymupdf as fitz
    try:
        doc = fitz.open(str(path))
        text = " ".join(doc[i].get_text() for i in range(min(3, doc.page_count)))
        doc.close()
    except Exception as e:
        return "review", f"extract_failed:{e}"
    if len(text.strip()) < 40:
        return "review", "no_text_layer"
    import difflib

    tt = _norm(paper.title or "")
    tx = _norm(text)
    title_hit = tt and (
        tt[:60] in tx or difflib.SequenceMatcher(None, tt, tx[: len(tt) + 200]).ratio() > 0.55
    )
    doi_hit = paper.doi and paper.doi.lower() in (
        text.lower().replace("doi.org/", "").replace("doi:", "doi:") or ""
    )
    rep_hit = paper.report_number and paper.report_number.lower() in text.lower()
    if doi_hit or rep_hit or title_hit:
        return "ok", f"title_match={bool(title_hit)} doi={bool(doi_hit)} rep={bool(rep_hit)}"
    # weak partial: >60% of title content words present
    words = [w for w in tt.split() if len(w) > 3]
    frac = sum(1 for w in words if w in tx) / max(1, len(words))
    if frac >= 0.6:
        return "weak", f"word_frac={frac:.2f}"
    return "mismatch", f"word_frac={frac:.2f}"


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------

STATUSES = (
    "SUCCESS",
    "METADATA_ONLY",
    "NO_RELEVANT_RESULT",
    "NO_PUBLIC_PDF",
    "ACCESS_RESTRICTED",
    "SOURCE_BLOCKED",
    "DOWNLOAD_FAILED",
    "PDF_MISMATCH",
    "PDF_INVALID",
    "NEEDS_REVIEW",
)


class Runner:
    def __init__(self, args, terms: list[Term]):
        self.args = args
        self.terms = terms
        self.out = Path(args.output_dir)
        self.out.mkdir(parents=True, exist_ok=True)
        self.pdf_dir = self.out / "pdfs"
        self.pdf_dir.mkdir(exist_ok=True)
        self.store = Store(self.out / "state.sqlite")
        self.run_id = datetime.now().strftime("run-%Y%m%d-%H%M%S")
        self.run_dir = self.out / "runs" / self.run_id
        self.run_dir.mkdir(parents=True, exist_ok=True)
        self.cfg = self._cfg()
        self.sources = self._build_sources()
        self.manager = None
        self.cache = None
        import threading

        self._mgr_lock = threading.Lock()
        self._log_lock = threading.Lock()
        self._search_log = open(  # noqa: SIM115 - held for the run's lifetime, closed in export
            self.run_dir / "search_attempts.jsonl", "a", encoding="utf-8"
        )
        self._dl_log = open(  # noqa: SIM115 - held for the run's lifetime, closed in export
            self.run_dir / "download_attempts.jsonl", "a", encoding="utf-8"
        )

    def _cfg(self) -> dict:
        db = self.out / "mosaic-cache.db"
        return {
            "db_path": str(db),
            "download_dir": str(self.pdf_dir),
            "filename_pattern": "{year}_{source}_{author}_{title}",
            "unpaywall": {"email": self.args.unpaywall_email},
            "download": {
                "backend": self.args.backend,
                "max_concurrent": self.args.max_concurrent,
                "per_host_concurrent": 3,
                "per_host_delay": 1.0,
                "max_tries": 3,
                "verify_pdf": True,
                "skip_existing": True,
                "aria2_rpc_url": "http://127.0.0.1:6800/jsonrpc",
                "aria2_rpc_secret": "",
                "aria2_auto_start": True,
                "aria2_fallback_native": True,
                "max_file_mb": self.args.max_file_mb,
                "max_split": 1,
                "max_connections": 1,
            },
        }

    def _build_sources(self):
        from mosaic.sources.arxiv import ArxivSource
        from mosaic.sources.base_search import BASESource
        from mosaic.sources.crossref import CrossrefSource
        from mosaic.sources.dblp import DBLPSource
        from mosaic.sources.doaj import DoajSource
        from mosaic.sources.hal import HALSource
        from mosaic.sources.ntrs import NTRSSource
        from mosaic.sources.openalex import OpenAlexSource
        from mosaic.sources.semantic_scholar import SemanticScholarSource
        from mosaic.sources.zenodo import ZenodoSource

        email = self.args.unpaywall_email
        wanted = set(self.args.sources.split(","))
        pool = {
            "openalex": OpenAlexSource(email=email),
            "crossref": CrossrefSource(email=email),
            "arxiv": ArxivSource(delay=3.0),
            "hal": HALSource(),
            "zenodo": ZenodoSource(),
            "doaj": DoajSource(),
            "base": BASESource(),
            "semantic_scholar": SemanticScholarSource(api_key=""),
            "ntrs": NTRSSource(delay=1.0),
            "dblp": DBLPSource(),
        }
        return {k: v for k, v in pool.items() if k in wanted}

    def _manager(self):
        with self._mgr_lock:
            if self.manager is None:
                from mosaic.db import Cache
                from mosaic.dlqueue import get_manager

                self.cache = Cache(self.cfg["db_path"])
                self.manager = get_manager(self.cfg, self.cache)
        return self.manager

    # -- search ----------------------------------------------------------

    def _search_one(self, source, query: str, n: int):
        return source.search(query, max_results=n, filters=None)

    def _try_source(self, name: str, source, query: str, n: int):
        """One source call for the parallel fan-out; health handled by caller."""
        if not self.store.source_ok(name):
            return name, [], "SKIPPED_COOLDOWN", "", 0.0
        t0 = time.time()
        try:
            papers = self._search_one(source, query, n)
            return name, papers, "ok", "", time.time() - t0
        except Exception as e:
            return name, [], "error", f"{type(e).__name__}: {e}", time.time() - t0

    def search_term(self, term: Term) -> tuple[list, int]:
        """Run tiered queries until budget exhausted or enough results."""
        from concurrent.futures import ThreadPoolExecutor

        from mosaic.services import merge_papers

        queries = plan_queries(term, self.args.max_queries_per_row)
        seen: dict[str, object] = {}
        ident_index: dict[str, str] = {}
        raw = 0
        stop = False
        for _qi, (query, tier) in enumerate(queries):
            if stop:
                break
            # Sources fan out in parallel *within* one query — each source is
            # still hit sequentially across queries, preserving per-source
            # politeness.
            active = {name: src for name, src in self.sources.items() if self.store.source_ok(name)}
            skipped = set(self.sources) - set(active)
            for name in skipped:
                self.store.log_search(term.term_id, query, tier, name, "SKIPPED_COOLDOWN", 0)
            with ThreadPoolExecutor(max_workers=4) as ex:
                results = list(
                    ex.map(
                        lambda kv, q=query: self._try_source(
                            kv[0], kv[1], q, self.args.max_results_per_source
                        ),
                        active.items(),
                    )
                )
            for name, papers, status, err, elapsed in results:
                if status == "ok":
                    self.store.source_success(name)
                elif status == "error":
                    cooldown = 3600.0 if "403" in err else (1200.0 if "429" in err else 0)
                    self.store.source_fail(name, err, cooldown=cooldown)
                self.store.log_search(term.term_id, query, tier, name, status, len(papers), err)
                with self._log_lock:
                    self._search_log.write(
                        json.dumps(
                            {
                                "term": term.term_id,
                                "query": query,
                                "tier": tier,
                                "source": name,
                                "status": status,
                                "n": len(papers),
                                "error": err,
                                "elapsed": round(elapsed, 2),
                                "ts": time.time(),
                            },
                            ensure_ascii=False,
                        )
                        + "\n"
                    )
                    self._search_log.flush()
                raw += len(papers)
                for p in papers:
                    self.store.save_paper(p)
                    merge_papers(seen, p, ident_index)
            # Early stop: enough HIGH-graded candidates after first two tiers
            if tier in ("L1_focus", "L2_ext") and len(seen) >= 4:
                grades = [score_paper(term, p)[0] for p in seen.values()]
                if grades.count("HIGH") >= 2:
                    stop = True
        return list(seen.values()), raw

    # -- download --------------------------------------------------------

    def _disk_ok(self) -> bool:
        used = sum(f.stat().st_size for f in self.pdf_dir.rglob("*.pdf"))
        if used >= self.args.max_disk_gb * 1024**3:
            return False
        free = shutil.disk_usage(self.pdf_dir).free
        return free > 1 * 1024**3

    def download_for_term(self, term: Term, ranked: list[tuple]) -> tuple[int, int]:
        """Enqueue top candidates; returns (download_ok, pdf_verified_ok)."""

        mgr = self._manager()
        ok_pdfs = 0
        dl_ok = 0
        # Select up to max_dl_candidates eligible papers, enqueue them as ONE
        # batch so aria2 fetches in parallel (queue max_concurrent), then
        # wait once. Sequential per-paper waits wasted a whole connect-timeout
        # on every bad candidate.
        batch: list[tuple] = []
        tried_uids: set[str] = set()
        for p, grade, _s in ranked:
            if len(batch) >= self.args.max_dl_candidates:
                break
            if p.uid in tried_uids:
                continue
            tried_uids.add(p.uid)
            if not (p.doi or p.pdf_url or p.pdf_candidates or p.ntrs_id):
                self.store.log_download(
                    term.term_id,
                    p.uid,
                    self.args.backend,
                    "",
                    "failed",
                    "no_candidates: no DOI/URL/attachment identifier",
                    "",
                    "",
                )
                continue
            has_path = any(
                (row["local_path"] or "") and Path(row["local_path"]).exists()
                for row in self.store.q("SELECT local_path FROM papers WHERE uid=?", (p.uid,))
            )
            if has_path:
                # Verified earlier (possibly for another term) — the same
                # PDF legitimately satisfies this term's link too.
                row = self.store.q("SELECT local_path,sha256 FROM papers WHERE uid=?", (p.uid,))[0]
                ok_pdfs += 1
                self.store.link(term.term_id, p.uid, grade, _s, True)
                self.store.log_download(
                    term.term_id,
                    p.uid,
                    self.args.backend,
                    p.pdf_url or "",
                    "ok",
                    "already_downloaded",
                    row["local_path"],
                    row["sha256"] or "",
                )
                continue
            if not self._disk_ok():
                log.warning("disk budget exhausted — pausing downloads")
                break
            batch.append((p, grade, _s))
        if batch:
            job_ids = mgr.enqueue_papers([p for p, _g, _s in batch])
            # enqueue_papers silently skips already-downloaded papers, so
            # resolve each job's paper via the job record — never zip.
            by_uid = {p.uid: (p, g, s) for p, g, s in batch}
            uid_for = {}
            for jid in job_ids:
                j = mgr.cache.get_job(jid)
                rec = by_uid.get(j["paper_uid"]) if j else None
                if rec:
                    uid_for[jid] = rec
            results = mgr.wait_all(job_ids, poll=0.4, timeout=self.args.term_dl_budget)
            # Papers silently skipped by enqueue (already verified in the
            # MOSAIC cache) still satisfy this term's requirement.
            queued_uids = {rec[0].uid for rec in uid_for.values()}
            for p, grade, _s in batch:
                if p.uid in queued_uids:
                    continue
                rec = mgr.cache.get_download(p.uid)
                ok_cached = bool(
                    rec
                    and rec["status"] == "ok"
                    and rec["local_path"]
                    and Path(rec["local_path"]).exists()
                )
                self.store.link(term.term_id, p.uid, grade, _s, ok_cached)
                self.store.log_download(
                    term.term_id,
                    p.uid,
                    self.args.backend,
                    p.pdf_url or "",
                    "ok" if ok_cached else "failed",
                    "already_downloaded" if ok_cached else "not enqueued / no cached file",
                    rec["local_path"] if ok_cached else "",
                    "",
                )
                if ok_cached:
                    self.store.save_paper(p, local_path=rec["local_path"], verify="ok")
                    ok_pdfs += 1
            cancelled: set[str] = set()
            for jid, (p, grade, _s) in uid_for.items():
                j = mgr.cache.get_job(jid)
                path = results.get(p.uid)
                if not path and j and j["status"] not in ("COMPLETED", "FAILED", "CANCELLED"):
                    mgr.cancel(jid)
                    cancelled.add(jid)
                if path:
                    dl_ok += 1
                    from mosaic.pdfcheck import sha256_file

                    sha = sha256_file(path)
                    ident, detail = pdf_identity(Path(path), p)
                    self.store.save_paper(
                        p, local_path=path, sha256=sha, verify="ok", identity=ident
                    )
                    self.store.log_download(
                        term.term_id,
                        p.uid,
                        self.args.backend,
                        p.pdf_url or "",
                        "ok" if ident != "mismatch" else "pdf_mismatch",
                        detail,
                        path,
                        sha,
                    )
                    with self._log_lock:
                        self._dl_log.write(
                            json.dumps(
                                {
                                    "term": term.term_id,
                                    "uid": p.uid,
                                    "result": "ok" if ident != "mismatch" else "pdf_mismatch",
                                    "identity": ident,
                                    "detail": detail,
                                    "path": path,
                                    "sha256": sha,
                                },
                                ensure_ascii=False,
                            )
                            + "\n"
                        )
                        self._dl_log.flush()
                    if ident != "mismatch":
                        ok_pdfs += 1
                        self.store.link(term.term_id, p.uid, grade, _s, True)
                else:
                    err = (j or {}).get("last_error") or "unknown"
                    if jid in cancelled:
                        err = f"timeout+{err}"
                    self.store.log_download(
                        term.term_id,
                        p.uid,
                        self.args.backend,
                        (j or {}).get("url") or p.pdf_url or "",
                        "failed",
                        err,
                        "",
                        "",
                    )
                    with self._log_lock:
                        self._dl_log.write(
                            json.dumps(
                                {
                                    "term": term.term_id,
                                    "uid": p.uid,
                                    "result": "failed",
                                    "error": err,
                                },
                                ensure_ascii=False,
                            )
                            + "\n"
                        )
                        self._dl_log.flush()
                    self.store.link(term.term_id, p.uid, grade, _s, False)
        return dl_ok, ok_pdfs

    # -- per-term pipeline ------------------------------------------------

    def process_term(self, term: Term, *, reuse_papers: bool = False) -> str:
        t0 = time.time()
        ranked: list[tuple] = []
        if reuse_papers:
            rows = self.store.relevant_links(term.term_id)
            for row in rows:
                pr = self.store.q("SELECT json FROM papers WHERE uid=?", (row["paper_uid"],))
                if pr:
                    from mosaic.models import Paper as _Paper

                    ranked.append(
                        (_Paper.from_dict(json.loads(pr[0]["json"])), row["grade"], row["score"])
                    )
        candidates = []
        raw_count = 0
        if not ranked:
            candidates, raw_count = self.search_term(term)
            scored = [(p, *score_paper(term, p)) for p in candidates]
            for p, grade, score, _why in scored:
                if grade in ("HIGH", "MEDIUM"):
                    self.store.link(term.term_id, p.uid, grade, score, False)
            accept = {"HIGH"} if term.domain not in BROAD_OK_DOMAINS else {"HIGH", "MEDIUM"}
            ranked = sorted(
                [x for x in scored if x[1] in accept],
                key=lambda x: (x[1] != "HIGH", not _dl_hint(x[0]), -x[2]),
            )
            ranked = [(p, g, s) for p, g, s, _w in ranked]
        else:
            candidates = ranked
        relevant = len(ranked)
        dl_ok = pdf_ok = 0
        if ranked and self._disk_ok():
            dl_ok, pdf_ok = self.download_for_term(term, ranked)
        status, reason, suggestion = self._classify(term, candidates, ranked, dl_ok, pdf_ok)
        self.store.save_term_result(
            term,
            status=status,
            queries=self.args.max_queries_per_row,
            candidates=raw_count or len(candidates),
            deduped=len(candidates),
            relevant=relevant,
            dl_candidates=min(len(ranked), self.args.max_dl_candidates),
            dl_ok=dl_ok,
            pdf_ok=pdf_ok,
            reason=reason,
            suggestion=suggestion,
        )
        log.info(
            "%s %-10s %s | cand=%d rel=%d dl=%d pdf=%d | %s (%.0fs)",
            term.term_id,
            status,
            term.en_core[:42],
            raw_count or len(candidates),
            relevant,
            dl_ok,
            pdf_ok,
            reason[:60],
            time.time() - t0,
        )
        return status

    def _classify(self, term, candidates, ranked, dl_ok, pdf_ok) -> tuple[str, str, str]:
        if pdf_ok > 0:
            return "SUCCESS", "", ""
        if not candidates:
            if any(not self.store.source_ok(n) for n in self.sources) and all(
                not self.store.source_ok(n) for n in self.sources
            ):
                return "SOURCE_BLOCKED", "all sources cooling down", "retry later"
            return (
                "NO_RELEVANT_RESULT",
                "no candidates from any source/tier",
                "try alternate spellings / CNKI manual",
            )
        if not ranked:
            return (
                "NO_RELEVANT_RESULT",
                f"{len(candidates)} candidates but none >= MEDIUM relevance",
                "review LOW/UNCERTAIN list; refine term alts",
            )
        # have relevant papers but no verified pdf
        rows = self.store.q(
            "SELECT reason FROM download_attempts WHERE term_id=? AND result='failed'",
            (term.term_id,),
        )
        reasons = {r["reason"].split(":")[0] for r in rows}
        texts = " ".join(r["reason"] for r in rows)
        if "no download URL" in texts or "no downloadable PDF candidates" in texts:
            return (
                "NO_PUBLIC_PDF",
                ";".join(sorted(reasons)),
                "check institutional access / author manuscript",
            )
        if reasons & {"access_denied"}:
            return (
                "ACCESS_RESTRICTED",
                ";".join(sorted(reasons)),
                "find OA copy via author/institution repo",
            )
        if reasons & {"not_found", "unreachable", "timeout", "server_error"}:
            return "DOWNLOAD_FAILED", ";".join(sorted(reasons)), "retry with --only-failed"
        if reasons & {"html_response", "invalid_pdf", "not_pdf", "missing_eof"}:
            return "PDF_INVALID", ";".join(sorted(reasons)), "alternate OA location needed"
        has_doi = any(p.doi for p, _g, _s in ranked)
        if has_doi:
            return (
                "NO_PUBLIC_PDF",
                "relevant papers found; no OA location resolved",
                "check institutional access / author manuscript",
            )
        return (
            "METADATA_ONLY",
            "relevant records but no DOI/URL resolved",
            "try NTRS/CNKI manual retrieval",
        )

    # -- export ------------------------------------------------------------

    def export(self) -> Path:
        out = self.run_dir
        for name, sql in {
            "term_results.csv": "SELECT * FROM term_results ORDER BY term_id",
            "papers.csv": "SELECT uid,title,year,doi,ntrs_id,source,pdf_url,local_path,sha256,verify_status,identity_status FROM papers ORDER BY uid",
            "term_paper_links.csv": "SELECT * FROM links ORDER BY term_id,grade,score DESC",
            "failures.csv": "SELECT * FROM term_results WHERE status!='SUCCESS' ORDER BY term_id",
        }.items():
            with open(out / name, "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow([d[0] for d in self.store.con.execute(sql).description])
                w.writerows(self.store.con.execute(sql))
        self._write_summary(out / "summary.md")
        return out

    def _write_summary(self, path: Path) -> None:
        rows = self.store.q("SELECT status, COUNT(*) c FROM term_results GROUP BY status")
        cnt = {r["status"]: r["c"] for r in rows}
        total = sum(cnt.values())
        uniq_papers = self.store.q("SELECT COUNT(*) c FROM papers")[0]["c"]
        uniq_pdfs = self.store.q(
            "SELECT COUNT(DISTINCT local_path) c FROM papers WHERE local_path!=''"
        )[0]["c"]
        terms_dl = self.store.q("SELECT COUNT(*) c FROM term_results WHERE pdf_ok>0")[0]["c"]
        lines = [
            f"# Validation run {self.run_id}",
            "",
            f"- terms executed: {total}/321",
            f"- terms with verified PDF: {terms_dl}",
            f"- unique papers recorded: {uniq_papers}",
            f"- unique verified PDFs: {uniq_pdfs}",
            "",
            "## Status breakdown",
        ]
        for s in STATUSES:
            if cnt.get(s):
                lines.append(f"- {s}: {cnt[s]}")
        path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def git_sha() -> str:
    try:
        return subprocess.run(
            ["git", "rev-parse", "HEAD"], capture_output=True, text=True, cwd=REPO
        ).stdout.strip()
    except Exception:
        return "unknown"


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", default="起步采集清单")
    ap.add_argument("--output-dir", default="./local-validation")
    ap.add_argument("--resume", action="store_true", help="skip terms already SUCCESS")
    ap.add_argument("--only-failed", action="store_true", help="retry non-SUCCESS terms")
    ap.add_argument("--start-id", default="")
    ap.add_argument("--end-id", default="")
    ap.add_argument("--ids", default="", help="comma list e.g. T002,T036")
    ap.add_argument("--sample", type=int, default=0, help="stratified sample of N terms")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--max-pdfs-per-row", type=int, default=1)
    ap.add_argument("--max-queries-per-row", type=int, default=6)
    ap.add_argument("--max-results-per-source", type=int, default=8)
    ap.add_argument("--max-dl-candidates", type=int, default=6)
    ap.add_argument(
        "--dl-timeout", type=int, default=600, help="per-paper download wall-clock limit (s)"
    )
    ap.add_argument(
        "--term-dl-budget", type=int, default=900, help="max seconds of download attempts per term"
    )
    ap.add_argument("--max-disk-gb", type=float, default=20.0)
    ap.add_argument("--max-file-mb", type=int, default=300)
    ap.add_argument("--max-concurrent", type=int, default=3)
    ap.add_argument("--workers", type=int, default=1, help="parallel term workers")
    ap.add_argument("--backend", default="aria2", choices=["aria2", "native", "auto"])
    ap.add_argument(
        "--sources",
        default="openalex,crossref,arxiv,hal,zenodo,doaj,base,semantic_scholar,ntrs,dblp",
    )
    ap.add_argument("--unpaywall-email", default="22371360@buaa.edu.cn")
    ap.add_argument("--export-only", action="store_true", help="just export CSVs from state")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
    )
    terms = load_terms(Path(args.xlsx))
    if not terms:
        log.error("no terms parsed")
        return 2

    # selection
    sel = terms
    if args.ids:
        want = {x.strip() for x in args.ids.split(",")}
        sel = [t for t in terms if t.term_id in want]
    if args.start_id:
        sel = [t for t in sel if t.term_id >= args.start_id]
    if args.end_id:
        sel = [t for t in sel if t.term_id <= args.end_id]
    if args.sample:
        per_domain: dict[str, list[Term]] = {}
        for t in sel:
            per_domain.setdefault(t.domain, []).append(t)
        sel = []
        quota = max(1, args.sample // max(1, len(per_domain)))
        for _d, ts in per_domain.items():
            sel.extend(ts[:quota])
        sel = sel[: args.sample]
        sel.sort(key=lambda t: (BATCH_ORDER.get(t.batch, 9), t.term_id))

    if args.dry_run:
        for t in sel:
            print(f"{t.term_id} [{t.batch}/{t.domain}] zh={t.zh_core} en={t.en_core}")
            for q, tier in plan_queries(t, args.max_queries_per_row):
                print(f"    {tier:11s} {q}")
        print(f"\n{len(sel)} terms selected")
        return 0

    runner = Runner(args, sel)
    runner.store.x(
        "INSERT OR REPLACE INTO run_meta(k,v) VALUES('run_id',?),('git_sha',?),('started',?)",
        (runner.run_id, git_sha(), datetime.now(UTC).isoformat()),
    )
    if args.export_only:
        print(runner.export())
        return 0

    counts: dict[str, int] = {}
    counts_lock = threading.Lock()

    def _run_one(term):
        st = runner.store.term_status(term.term_id)
        if args.resume and st == "SUCCESS":
            return st, False
        if args.only_failed and st == "SUCCESS":
            return None, False
        # Reuse stored links for any previously-attempted-but-failed term:
        # search evidence is kept, only the download stage is retried.
        # Terms with no relevant links still get a fresh search below.
        reuse = bool(st and st != "SUCCESS")
        try:
            return runner.process_term(term, reuse_papers=reuse), True
        except KeyboardInterrupt:
            raise
        except Exception as e:
            log.exception("term %s crashed", term.term_id)
            runner.store.save_term_result(term, status="NEEDS_REVIEW", reason=f"harness error: {e}")
            return "NEEDS_REVIEW", True

    todo = []
    for term in sel:
        st = runner.store.term_status(term.term_id)
        if args.resume and st == "SUCCESS":
            counts[st] = counts.get(st, 0) + 1
        elif args.only_failed and st == "SUCCESS":
            continue
        else:
            todo.append(term)

    try:
        if args.workers > 1:
            from concurrent.futures import ThreadPoolExecutor, as_completed

            with ThreadPoolExecutor(max_workers=args.workers) as ex:
                futs = {ex.submit(_run_one, t): t for t in todo}
                for fut in as_completed(futs):
                    status, ran = fut.result()
                    if ran and status:
                        with counts_lock:
                            counts[status] = counts.get(status, 0) + 1
        else:
            for term in todo:
                status, ran = _run_one(term)
                if ran and status:
                    counts[status] = counts.get(status, 0) + 1
    finally:
        out = runner.export()
        log.info("run dir: %s", out)
        log.info("status counts: %s", counts)
        from mosaic.dlqueue import shutdown_managers

        shutdown_managers()
    return 0


if __name__ == "__main__":
    sys.exit(main())
